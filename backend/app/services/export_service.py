"""Export Service - Generate PDF and Excel exports of comparison data."""

import io
import logging
from typing import Optional
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

logger = logging.getLogger(__name__)


class ExportService:
    """Service for exporting comparison data to PDF and Excel."""

    @staticmethod
    def export_excel(
        project_name: str,
        vendors: list[dict],
        comparison_results: list[dict],
        summary: Optional[dict] = None
    ) -> bytes:
        """Generate Excel export of comparison data."""
        wb = Workbook()

        # ---- Comparison Sheet ----
        ws = wb.active
        ws.title = "Vendor Comparison"

        # Header styling
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="1a1a2e", end_color="1a1a2e", fill_type="solid")
        lowest_fill = PatternFill(start_color="d4edda", end_color="d4edda", fill_type="solid")
        highest_fill = PatternFill(start_color="f8d7da", end_color="f8d7da", fill_type="solid")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # Title
        ws.merge_cells('A1:F1')
        title_cell = ws['A1']
        title_cell.value = f"Vendor Comparison Report - {project_name}"
        title_cell.font = Font(bold=True, size=14, color="1a1a2e")
        title_cell.alignment = Alignment(horizontal='center')

        # Build headers
        vendor_names = [v.get("name", f"Vendor {i+1}") for i, v in enumerate(vendors)]
        headers = ["#", "Item Name", "Qty", "Unit"]
        for vname in vendor_names:
            headers.extend([f"{vname} (Unit)", f"{vname} (Total)"])
        headers.extend(["Lowest Vendor", "Price Diff %", "Delivery", "Remarks"])

        row = 3
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', wrap_text=True)
            cell.border = border

        # Data rows
        for idx, result in enumerate(comparison_results, 1):
            row += 1
            col = 1

            ws.cell(row=row, column=col, value=idx).border = border
            col += 1
            ws.cell(row=row, column=col, value=result.get("normalized_item_name", "")).border = border
            col += 1

            # Get first vendor's quantity/unit
            first_vendor = next(iter(result.get("vendor_prices", {}).values()), {})
            ws.cell(row=row, column=col, value=first_vendor.get("quantity", "")).border = border
            col += 1
            ws.cell(row=row, column=col, value=first_vendor.get("unit", "")).border = border
            col += 1

            # Vendor prices
            lowest_vid = result.get("lowest_vendor_id")
            for vendor in vendors:
                vid = vendor["id"]
                vp = result.get("vendor_prices", {}).get(vid, {})

                unit_cell = ws.cell(row=row, column=col, value=vp.get("unit_price", "-"))
                unit_cell.border = border
                if vid == lowest_vid:
                    unit_cell.fill = lowest_fill
                col += 1

                total_cell = ws.cell(row=row, column=col, value=vp.get("total_price", "-"))
                total_cell.border = border
                if vid == lowest_vid:
                    total_cell.fill = lowest_fill
                col += 1

            # Lowest vendor name
            lowest_name = ""
            if lowest_vid:
                for v in vendors:
                    if v["id"] == lowest_vid:
                        lowest_name = v.get("name", "")
                        break

            ws.cell(row=row, column=col, value=lowest_name).border = border
            col += 1
            ws.cell(row=row, column=col, value=result.get("price_difference_pct", "")).border = border
            col += 1
            ws.cell(row=row, column=col, value=first_vendor.get("delivery_days", "")).border = border
            col += 1
            ws.cell(row=row, column=col, value=result.get("remarks", "")).border = border

        # Auto-adjust column widths
        for col in ws.columns:
            max_length = 0
            col_letter = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            ws.column_dimensions[col_letter].width = min(max_length + 4, 30)

        # ---- Summary Sheet ----
        if summary:
            ws2 = wb.create_sheet("Summary")
            ws2['A1'] = "Procurement Intelligence Summary"
            ws2['A1'].font = Font(bold=True, size=14)

            row = 3
            if "overall_assessment" in summary:
                ws2.cell(row=row, column=1, value="Overall Assessment").font = Font(bold=True)
                row += 1
                ws2.cell(row=row, column=1, value=summary["overall_assessment"])
                row += 2

            if "recommendations" in summary:
                ws2.cell(row=row, column=1, value="Recommendations").font = Font(bold=True)
                row += 1
                for rec in summary["recommendations"]:
                    ws2.cell(row=row, column=1, value=f"• {rec.get('recommendation', '')}")
                    row += 1

            if "summary_text" in summary:
                row += 1
                ws2.cell(row=row, column=1, value="Executive Summary").font = Font(bold=True)
                row += 1
                ws2.cell(row=row, column=1, value=summary["summary_text"])

        # Save to bytes
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    @staticmethod
    def export_pdf(
        project_name: str,
        vendors: list[dict],
        comparison_results: list[dict],
        summary: Optional[dict] = None
    ) -> bytes:
        """Generate PDF export of comparison data."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=landscape(A4),
            rightMargin=30,
            leftMargin=30,
            topMargin=30,
            bottomMargin=30
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Title'],
            fontSize=18,
            textColor=colors.HexColor("#1a1a2e"),
            spaceAfter=20
        )
        subtitle_style = ParagraphStyle(
            'Subtitle',
            parent=styles['Heading2'],
            fontSize=13,
            textColor=colors.HexColor("#16213e"),
            spaceAfter=10
        )

        elements = []

        # Title
        elements.append(Paragraph(f"Vendor Comparison Report", title_style))
        elements.append(Paragraph(f"Project: {project_name}", subtitle_style))
        elements.append(Spacer(1, 12))

        # Build comparison table
        vendor_names = [v.get("name", f"Vendor {i+1}") for i, v in enumerate(vendors)]
        headers = ["#", "Item", "Qty"]
        for vname in vendor_names:
            headers.append(f"{vname}\nPrice")
        headers.extend(["Lowest", "Diff %"])

        table_data = [headers]

        for idx, result in enumerate(comparison_results, 1):
            row = [str(idx)]
            row.append(result.get("normalized_item_name", "")[:30])
            first_vendor = next(iter(result.get("vendor_prices", {}).values()), {})
            row.append(str(first_vendor.get("quantity", "-")))

            for vendor in vendors:
                vid = vendor["id"]
                vp = result.get("vendor_prices", {}).get(vid, {})
                price = vp.get("unit_price") or vp.get("total_price") or "-"
                row.append(str(price))

            lowest_name = ""
            lowest_vid = result.get("lowest_vendor_id")
            if lowest_vid:
                for v in vendors:
                    if v["id"] == lowest_vid:
                        lowest_name = v.get("name", "")[:15]
                        break
            row.append(lowest_name)
            row.append(str(result.get("price_difference_pct", "-")))

            table_data.append(row)

        # Create table with styling
        col_count = len(headers)
        col_widths = [30, 120, 40] + [70] * len(vendors) + [80, 50]

        table = Table(table_data, colWidths=col_widths[:col_count])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1a1a2e")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 8),
            ('FONTSIZE', (0, 1), (-1, -1), 7),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
            ('BACKGROUND', (0, 1), (-1, -1), colors.white),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8f9fa")]),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ]))

        elements.append(table)

        # Add summary if available
        if summary:
            elements.append(Spacer(1, 20))
            elements.append(Paragraph("AI Procurement Summary", subtitle_style))

            if "overall_assessment" in summary:
                elements.append(Paragraph(summary["overall_assessment"], styles["Normal"]))
                elements.append(Spacer(1, 8))

            if "summary_text" in summary:
                elements.append(Paragraph(summary["summary_text"], styles["Normal"]))

        doc.build(elements)
        buffer.seek(0)
        return buffer.getvalue()
